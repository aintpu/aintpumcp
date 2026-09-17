from __future__ import annotations

import hashlib
import json
import os
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

import openpyxl


OFFICES = {
    "ope": "體育室",
    "ge": "通識教育中心",
    "lc": "語言中心",
    "oaa": "教務處",
    "osa": "學務處",
    "hr": "人事室",
    "oga": "總務處",
}

OFFICE_HOME_URLS = {
    "ope": "https://new.ntpu.edu.tw/ope",
    "ge": "https://new.ntpu.edu.tw/cge",
    "lc": "https://lc.ntpu.edu.tw",
    "oaa": "https://new.ntpu.edu.tw/oaa",
    "osa": "https://new.ntpu.edu.tw/osa",
    "hr": "https://new.ntpu.edu.tw/op",
    "oga": "https://new.ntpu.edu.tw/oga",
}

PORTAL_FILES = {
    "ope": "all_content_v2.md",
    "ge": "cge_content.md",
    "lc": "lc_content.md",
    "hr": "hr_content.md",
    "oga": "oga_content.md",
}

REGULATION_FILES = {
    "ope": "ALL_files_2.md",
    "oaa": "oaa_regulations.md",
    "osa": "osa_regulations.md",
}

ACADEMIC_SHEETS = {"ge": "通識教育中心", "lc": "語言中心"}
ADMIN_LABELS = {"oaa": "教務處", "osa": "學生事務處"}
DOCUMENT_BASE_URL = os.getenv(
    "NTPU_AIA_DOCUMENT_BASE_URL",
    "https://aia.mcp.ntpu.ai/documents",
).rstrip("/")

H1_RE = re.compile(r"^#\s+(.+?)\s*$", re.M)
H2_RE = re.compile(r"^##\s+(.+?)\s*$", re.M)
H3_RE = re.compile(r"^###\s+(.+?)\s*$", re.M)
PAGE_RE = re.compile(r"^###\s+(Page\s+[^\n]+)\s*$", re.M | re.I)
MARKDOWN_LINK_RE = re.compile(r"\[([^\]]+)\]\((https?://[^)\s]+)\)")
PLAIN_URL_RE = re.compile(r"https?://[^\s<>\])]+")
DATE_RE = re.compile(r"(20\d{2})\s*/\s*(\d{1,2})\s*/\s*(\d{1,2})")


@dataclass(frozen=True)
class KnowledgeDocument:
    source_id: str
    office_code: str
    office_name: str
    title: str
    content_type: str
    page: str | None
    published_at: str | None
    text: str
    source_url: str | None
    source_file: str
    source_commit: str | None
    verified_at: str
    content_hash: str


def _normalise_title(value: str) -> str:
    value = re.sub(r"\.(pdf|docx?|odt|ods)$", "", value.strip(), flags=re.I)
    return re.sub(r"[\s　]+", "", value)


def _clean_title(value: str) -> str:
    value = re.sub(r"^#+\s*", "", value.strip())
    match = MARKDOWN_LINK_RE.search(value)
    if match:
        value = match.group(1)
    return value.strip("*_ `")


def _first_url(text: str) -> str | None:
    explicit = re.findall(r"(?:來源網址|檔案連結|資料來源網址)\s*[:：]\s*(https?://\S+)", text)
    if explicit:
        return explicit[-1].rstrip(".,，。)")
    link = MARKDOWN_LINK_RE.search(text)
    if link:
        return link.group(2)
    plain = PLAIN_URL_RE.search(text)
    return plain.group(0) if plain else None


def _published_at(text: str) -> str | None:
    match = DATE_RE.search(text)
    if not match:
        return None
    year, month, day = (int(part) for part in match.groups())
    return f"{year:04d}-{month:02d}-{day:02d}"


def _public_source_url(url: str | None) -> str | None:
    old_prefix = "https://aia.ntpu.ai/documents/"
    if url and url.startswith(old_prefix):
        return f"{DOCUMENT_BASE_URL}/{url.removeprefix(old_prefix)}"
    return url


def _split_sections(text: str, pattern: re.Pattern[str]) -> list[tuple[str, str]]:
    matches = list(pattern.finditer(text))
    sections: list[tuple[str, str]] = []
    for index, match in enumerate(matches):
        end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
        sections.append((match.group(1).strip(), text[match.end():end].strip()))
    return sections


def _chunk_text(text: str, max_chars: int = 1800, overlap: int = 180) -> list[str]:
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if not text:
        return []
    if len(text) <= max_chars:
        return [text]

    paragraphs = [item.strip() for item in re.split(r"\n\s*\n", text) if item.strip()]
    chunks: list[str] = []
    current = ""
    for paragraph in paragraphs:
        parts = [paragraph[i:i + max_chars] for i in range(0, len(paragraph), max_chars)]
        for part in parts:
            candidate = f"{current}\n\n{part}".strip() if current else part
            if len(candidate) <= max_chars:
                current = candidate
                continue
            if current:
                chunks.append(current)
                prefix = current[-overlap:].lstrip()
                current = f"{prefix}\n\n{part}".strip()
            else:
                chunks.append(part)
                current = ""
    if current:
        chunks.append(current)
    return chunks


def _manifest(data_root: Path) -> dict:
    path = data_root / "manifest.json"
    if not path.exists():
        return {"source_commit": None, "generated_at": "unknown", "files": {}}
    return json.loads(path.read_text(encoding="utf-8"))


def _content_hash(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def _make_documents(
    *,
    office: str,
    title: str,
    text: str,
    content_type: str,
    source_file: str,
    manifest: dict,
    source_url: str | None = None,
    page: str | None = None,
    published_at: str | None = None,
) -> list[KnowledgeDocument]:
    documents: list[KnowledgeDocument] = []
    title = _clean_title(title) or "未命名文件"
    source_url = _public_source_url(
        source_url or _first_url(text) or OFFICE_HOME_URLS[office]
    )
    for chunk_index, chunk in enumerate(_chunk_text(text)):
        seed = f"{office}|{source_file}|{title}|{page or ''}|{chunk_index}"
        source_id = hashlib.sha256(seed.encode("utf-8")).hexdigest()[:24]
        documents.append(KnowledgeDocument(
            source_id=source_id,
            office_code=office,
            office_name=OFFICES[office],
            title=title,
            content_type=content_type,
            page=page,
            published_at=published_at,
            text=chunk,
            source_url=source_url,
            source_file=source_file,
            source_commit=manifest.get("source_commit"),
            verified_at=manifest.get("generated_at") or "unknown",
            content_hash=_content_hash(chunk),
        ))
    return documents


def _load_academic_metadata(crawler_dir: Path) -> dict[str, dict[str, dict[str, str]]]:
    path = crawler_dir / "北大學術單位法規彙整.xlsx"
    output: dict[str, dict[str, dict[str, str]]] = {code: {} for code in ACADEMIC_SHEETS}
    if not path.exists():
        return output
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        for office, sheet_name in ACADEMIC_SHEETS.items():
            if sheet_name not in workbook.sheetnames:
                continue
            rows = workbook[sheet_name].iter_rows(values_only=True)
            headers = [str(cell or "") for cell in next(rows)]
            for row in rows:
                item = dict(zip(headers, row))
                title = str(item.get("title") or "").strip()
                if not title:
                    continue
                output[office][_normalise_title(title)] = {
                    "url": str(item.get("file_url") or item.get("source_page") or "").strip(),
                    "date": str(item.get("updated_date") or "").strip(),
                    "type": "form" if "表單" in str(item.get("tags") or "") else "regulation",
                }
    finally:
        workbook.close()
    return output


def _load_admin_metadata(crawler_dir: Path) -> dict[str, dict[str, dict[str, str]]]:
    path = crawler_dir / "北大行政單位法規彙整.xlsx"
    output: dict[str, dict[str, dict[str, str]]] = {code: {} for code in ADMIN_LABELS}
    if not path.exists():
        return output
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        rows = workbook[workbook.sheetnames[0]].iter_rows(values_only=True)
        headers = [str(cell or "") for cell in next(rows)]
        for row in rows:
            item = dict(zip(headers, row))
            for office, label in ADMIN_LABELS.items():
                if str(item.get("處室") or "").strip() != label:
                    continue
                title = str(item.get("法規名稱") or "").strip()
                if not title:
                    continue
                output[office][_normalise_title(title)] = {
                    "url": str(item.get("檔案連結") or item.get("網址") or "").strip(),
                    "date": str(item.get("上傳日期") or "").strip(),
                    "type": "form" if "表單" in str(item.get("標籤") or "") else "regulation",
                }
    finally:
        workbook.close()
    return output


def _load_ope_regulation_metadata(text: str) -> dict[str, dict[str, str]]:
    output: dict[str, dict[str, str]] = {}
    for line in text.splitlines():
        if not line.startswith("- ") or " | " not in line:
            continue
        parts = [part.strip() for part in line[2:].split("|")]
        if len(parts) < 4:
            continue
        title, category, _, url = parts[:4]
        output[_normalise_title(title)] = {
            "url": url,
            "date": "",
            "type": "form" if category == "表單" else "regulation",
        }
    return output


def _parse_faq_section(
    office: str,
    body: str,
    source_file: str,
    manifest: dict,
) -> list[KnowledgeDocument]:
    documents: list[KnowledgeDocument] = []
    for title, answer in _split_sections(body, H3_RE):
        if title.lower().startswith(("page ", "公告內容", "相關下載")) or len(answer) < 10:
            continue
        documents.extend(_make_documents(
            office=office,
            title=title,
            text=answer,
            content_type="faq",
            source_file=source_file,
            manifest=manifest,
        ))
    return documents


def _parse_news_section(
    office: str,
    body: str,
    source_file: str,
    manifest: dict,
) -> list[KnowledgeDocument]:
    documents: list[KnowledgeDocument] = []
    blocks = re.split(r"^---\s*$", body, flags=re.M)
    for block in blocks:
        link = MARKDOWN_LINK_RE.search(block)
        if not link or len(block.strip()) < 40:
            continue
        title, url = link.groups()
        documents.extend(_make_documents(
            office=office,
            title=title,
            text=block,
            content_type="news",
            source_file=source_file,
            manifest=manifest,
            source_url=url,
            published_at=_published_at(block),
        ))
    return documents


def _parse_regulation_sections(
    office: str,
    body: str,
    source_file: str,
    manifest: dict,
    metadata: dict[str, dict[str, str]],
) -> list[KnowledgeDocument]:
    documents: list[KnowledgeDocument] = []
    for title, regulation_body in _split_sections(body, H2_RE):
        if len(regulation_body) < 20 or title in {"文件索引", "表單索引"}:
            continue
        meta = metadata.get(_normalise_title(title), {})
        pages = _split_sections(regulation_body, PAGE_RE)
        if not pages:
            pages = [("總覽", regulation_body)]
        for page, page_text in pages:
            documents.extend(_make_documents(
                office=office,
                title=title,
                text=page_text,
                content_type=meta.get("type", "regulation"),
                source_file=source_file,
                manifest=manifest,
                source_url=meta.get("url") or None,
                page=page,
                published_at=meta.get("date") or _published_at(regulation_body),
            ))
    return documents


def _parse_portal_file(
    office: str,
    path: Path,
    manifest: dict,
    academic_metadata: dict[str, dict[str, dict[str, str]]],
) -> list[KnowledgeDocument]:
    text = path.read_text(encoding="utf-8-sig")
    source_file = f"crawler_data/{path.name}"
    documents: list[KnowledgeDocument] = []
    for category, body in _split_sections(text, H1_RE):
        if "常見問題" in category:
            documents.extend(_parse_faq_section(office, body, source_file, manifest))
        elif "最新消息" in category:
            documents.extend(_parse_news_section(office, body, source_file, manifest))
        elif "### Page" in body:
            documents.extend(_parse_regulation_sections(
                office, body, source_file, manifest,
                academic_metadata.get(office, {}),
            ))
        else:
            sections = _split_sections(body, H2_RE) or [(category, body)]
            for title, section_body in sections:
                if len(section_body) < 20:
                    continue
                documents.extend(_make_documents(
                    office=office,
                    title=title,
                    text=section_body,
                    content_type="page",
                    source_file=source_file,
                    manifest=manifest,
                ))
    return documents


def load_documents(data_root: Path) -> list[KnowledgeDocument]:
    crawler_dir = data_root / "crawler_data"
    manifest = _manifest(data_root)
    academic_metadata = _load_academic_metadata(crawler_dir)
    admin_metadata = _load_admin_metadata(crawler_dir)
    documents: list[KnowledgeDocument] = []

    for office, filename in PORTAL_FILES.items():
        path = crawler_dir / filename
        if path.exists():
            documents.extend(_parse_portal_file(
                office, path, manifest, academic_metadata,
            ))

    for office, filename in REGULATION_FILES.items():
        path = crawler_dir / filename
        if not path.exists():
            continue
        text = path.read_text(encoding="utf-8-sig")
        metadata = (
            _load_ope_regulation_metadata(text)
            if office == "ope"
            else admin_metadata.get(office, {})
        )
        documents.extend(_parse_regulation_sections(
            office=office,
            body=text,
            source_file=f"crawler_data/{filename}",
            manifest=manifest,
            metadata=metadata,
        ))

    unique: dict[str, KnowledgeDocument] = {}
    for document in documents:
        unique[document.source_id] = document
    return list(unique.values())


def office_counts(documents: Iterable[KnowledgeDocument]) -> dict[str, int]:
    counts = {office: 0 for office in OFFICES}
    for document in documents:
        counts[document.office_code] += 1
    return counts
